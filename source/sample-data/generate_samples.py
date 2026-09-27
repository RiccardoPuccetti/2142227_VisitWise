"""Generates the import template and SYNTHETIC sample datasets (no real/NDA data).

Usage (from repo root):  python source/sample-data/generate_samples.py
Requires: pip install openpyxl

Outputs:
  source/backend/src/main/resources/import-template/visitwise-import-template.xlsx  (served by GET /api/imports/template)
  source/sample-data/sample-erp-layout.xlsx       same layout as the real ERP export (Italian headers, subtotal rows, '-' cells)
  source/sample-data/sample-shuffled-columns.xlsx same data, columns shuffled and renamed (tests the wizard's column mapping)
"""
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE.parent / "backend" / "src" / "main" / "resources" / "import-template" / "visitwise-import-template.xlsx"

# Real streets (so geocoding works), fake businesses.
ADDRESSES = [
    ("VIA DEL CORSO 300", "ROMA"), ("VIA NAZIONALE 75", "ROMA"), ("VIA APPIA NUOVA 210", "ROMA"),
    ("VIALE TRASTEVERE 60", "ROMA"), ("VIA COLA DI RIENZO 150", "ROMA"), ("VIA TUSCOLANA 700", "ROMA"),
    ("VIA OSTIENSE 95", "ROMA"), ("VIA NOMENTANA 250", "ROMA"), ("VIA CASILINA 400", "ROMA"),
    ("VIA SALARIA 220", "ROMA"), ("VIALE MARCONI 180", "ROMA"), ("VIA PRENESTINA 300", "ROMA"),
    ("VIA DEI CONDOTTI 20", "ROMA"), ("PIAZZA NAVONA 45", "ROMA"), ("VIA VENETO 120", "ROMA"),
    ("VIA MERULANA 140", "ROMA"), ("VIA LABICANA 90", "ROMA"), ("VIA FLAMINIA 350", "ROMA"),
    ("VIA AURELIA 480", "ROMA"), ("VIA CASSIA 600", "ROMA"), ("VIALE EUROPA 100", "ROMA"),
    ("VIA LIBERIANA 17", "ROMA"), ("VIA DEL TRITONE 60", "ROMA"), ("VIA GREGORIO VII 200", "ROMA"),
    ("VIALE REGINA MARGHERITA 190", "ROMA"), ("VIA TIBURTINA 500", "ROMA"), ("VIA TUSCOLANA 1200", "ROMA"),
    ("VIA DI PORTA MAGGIORE 30", "ROMA"), ("PIAZZA DEL POPOLO 18", "ROMA"), ("VIA CAVOUR 250", "ROMA"),
    ("PIAZZA GARIBALDI 10", "TIVOLI"), ("CORSO MATTEOTTI 50", "ALBANO LAZIALE"),
    ("PIAZZA DELLA REPUBBLICA 5", "MARINO"), ("VIA DEI CALZAIUOLI 30", "FIRENZE"),
]
PREFIXES = ["OSTERIA", "TRATTORIA", "ENOTECA", "BAR", "RISTORANTE", "BISTROT", "WINE BAR", "HOTEL"]
NAMES = ["AURORA", "DEL PORTO", "LA PERGOLA", "IL GLICINE", "SAN MARCO", "LE TERRAZZE", "DA NINO",
         "IL CORTILE", "ARCOBALENO", "LA LANTERNA", "DEI FIORI", "IL GIARDINO", "LUNA ROSSA",
         "IL FARO", "DEL SOLE", "BELVEDERE", "LA BOTTEGA", "IL CAMINETTO"]
AGENTS = ["AGENT NORTH", "AGENT SOUTH", "AGENT EAST", "AGENT WEST"]
ENTERPRISES = ["ENTERPRISE A", "ENTERPRISE B", "ENTERPRISE C"]


def build_rows(seed: int = 42):
    rnd = random.Random(seed)
    rows = []
    addr_iter = iter(ADDRESSES)
    customer_id = 70000
    for _ in range(55):
        customer_id += rnd.randint(10, 900)
        company = f"{rnd.choice(NAMES)} {rnd.choice(['SRL', 'SNC', 'SAS', 'SRLS'])}"
        customer = f"{company} ({customer_id})"
        for _ in range(rnd.choice([1, 1, 1, 2])):
            try:
                address, city = next(addr_iter)
            except StopIteration:
                # Same real street, another house number.
                street, city = rnd.choice(ADDRESSES[:30])
                address = f"{street.rsplit(' ', 1)[0]} {rnd.randint(2, 200)}"
            point = f"{rnd.choice(PREFIXES)} {rnd.choice(NAMES)}"
            agent = rnd.choice(AGENTS)
            values = []
            for _ in ENTERPRISES:
                if rnd.random() < 0.45:
                    values.append(round(rnd.lognormvariate(7.5, 1.3), 2))
                else:
                    values.append("-")
            if all(v == "-" for v in values):
                values[0] = round(rnd.lognormvariate(7, 1), 2)
            if rnd.random() < 0.05:  # credit note
                values[0] = -round(rnd.uniform(50, 300), 2)
            rows.append((customer, point, address, city, agent, values))
    return rows


def num(v):
    return 0 if v == "-" else v


def write_erp_layout(rows, path: Path):
    """Mimics an ERP export: Italian headers, 'Totale' subtotal row after each customer, grand total."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Ragione Sociale", "Punto Vendita", "Indirizzo", "Comune",
               "Agente", "Totale", *ENTERPRISES])
    grand = [0.0] * (len(ENTERPRISES) + 1)
    by_customer: dict[str, list] = {}
    for r in rows:
        by_customer.setdefault(r[0], []).append(r)
    for customer, crow in by_customer.items():
        sub = [0.0] * (len(ENTERPRISES) + 1)
        for (_, point, address, city, agent, values) in crow:
            total = round(sum(num(v) for v in values), 2)
            ws.append([customer, point, address, city, agent, total, *values])
            sub[0] += total
            for i, v in enumerate(values):
                sub[i + 1] += num(v)
        ws.append([f"{customer} Totale", None, None, None, None, *[round(x, 2) for x in sub]])
        grand = [g + s for g, s in zip(grand, sub)]
    ws.append(["Totale complessivo", None, None, None, None, *[round(x, 2) for x in grand]])
    wb.save(path)


def write_shuffled(rows, path: Path):
    """Same detail rows, different column order and English headers, no subtotal rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    ws.append(["ENTERPRISE C", "City", "Sales Agent", "Customer", "Delivery Address", "ENTERPRISE A",
               "Point of Sale", "ENTERPRISE B"])
    for (customer, point, address, city, agent, values) in rows:
        w, v, s = values
        ws.append([s, city, agent, customer, address, w, point, v])
    wb.save(path)


def write_template(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    headers = ["Customer", "Delivery Point", "Address", "City", "Agent", "Latitude", "Longitude",
               "Enterprise A", "Enterprise B", "Enterprise C"]
    ws.append(headers)
    ws.append(["ACME SRL (10001)", "RISTORANTE ACME", "VIA DEL CORSO 300", "ROMA", "AGENT NORTH",
               None, None, 1250.50, "-", 830])
    ws.append(["ACME SRL (10001)", "BAR ACME", "VIA NAZIONALE 75", "ROMA", "AGENT NORTH",
               41.9009, 12.4923, "-", 420, "-"])
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1E3A8A")
    for cell in ws[1]:
        cell.font = bold
        cell.fill = fill
    for col, width in zip("ABCDEFGHIJ", [28, 26, 30, 16, 16, 11, 11, 14, 14, 14]):
        ws.column_dimensions[col].width = width

    info = wb.create_sheet("Instructions")
    for line in [
        ["VisitWise import template"],
        [],
        ["Required columns", "Customer, Delivery Point, Address, City"],
        ["Optional columns", "Agent, Latitude, Longitude (if present, geocoding is skipped for that row)"],
        ["Enterprise columns", "One column per federation company with the yearly revenue (any number, any name)."],
        ["", "Rename 'Enterprise A/B/C' to the names of your companies."],
        ["Empty values", "Empty cells or '-' are read as 0. Negative values (credit notes) are allowed."],
        ["Column order", "Free: in the import wizard you map each column to its meaning."],
        ["Subtotal rows", "Rows without Delivery Point / Address (e.g. '... Totale') are skipped automatically."],
    ]:
        info.append(line)
    info["A1"].font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 90
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


if __name__ == "__main__":
    data = build_rows()
    write_template(TEMPLATE)
    write_erp_layout(data, HERE / "sample-erp-layout.xlsx")
    write_shuffled(data, HERE / "sample-shuffled-columns.xlsx")
    print(f"{len(data)} synthetic rows written. Template: {TEMPLATE}")
